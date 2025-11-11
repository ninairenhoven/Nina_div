import pandas as pd
from docx import Document
import os
from pathlib import Path
from datetime import datetime
import logging
import shutil
import gc
from xhtml2pdf import pisa
from concurrent.futures import ProcessPoolExecutor, as_completed
import multiprocessing
import time
import re
from tqdm import tqdm  # Progress bar

# Mammoth import flyttet til toppen for å unngå gjentatt import
try:
    from mammoth import convert_to_html
    MAMMOTH_AVAILABLE = True
except ImportError:
    MAMMOTH_AVAILABLE = False
    print("Warning: mammoth not installed, using basic conversion")


# =============================================================================
# INPUT FILES
# =============================================================================

user_path = Path.home()
path = user_path.joinpath(r'Documents\RVU_LOKAL\BRAKAR Sample')
TEMPLATE_FILE = path.joinpath("Invitasjonsbrev_ORDA_Brakar.docx")
LOOKUP_FILE = path.joinpath("Buskerud FK, Brakar, Opinion FREG uttrekk-krr_BRAKAR_Batch_0_23_PROCESSED_Nov11.csv")

timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
OUTPUT_FOLDER = path.joinpath(f"OUTPUT_{timestamp}")

# VIKTIG: For beste ytelse, bruk SSD for OUTPUT_FOLDER!
# OUTPUT_FOLDER = Path("D:/FastDisk/brakar_output")  # Eksempel: Dedikert SSD


# =============================================================================
# FONT CONFIGURATION
# =============================================================================
DEFAULT_FONT_FAMILY = "Helvetica, Arial, sans-serif"
DEFAULT_FONT_SIZE = "11pt"
DEFAULT_LINE_HEIGHT = "1.5"
LIST_LINE_HEIGHT = "1.2"
HEADING_1_SIZE = "14pt"
HEADING_2_SIZE = "12pt"


# =============================================================================
# GLOBAL WORKER TEMPLATE CACHE - Optimalisering for multiprocessing
# =============================================================================
_worker_template_cache = None


def init_worker(template_path):
    """
    Initialize worker process with template (called once per worker).
    Dette sparer minnebruk ved å unngå å sende template via pickle til hver worker.
    """
    global _worker_template_cache
    if _worker_template_cache is None:
        _worker_template_cache = load_template_for_worker(template_path)


def load_template_for_worker(template_path):
    """Load and convert template to HTML in worker process"""
    try:
        if MAMMOTH_AVAILABLE:
            with open(template_path, "rb") as docx_file:
                result = convert_to_html(docx_file)
                html_body = result.value
        else:
            doc = Document(template_path)
            paragraphs = []
            for para in doc.paragraphs:
                text = para.text
                if text.strip():
                    paragraphs.append(f"<p>{text}</p>")
            html_body = "\n".join(paragraphs)
        
        # Process spacing tags and emails
        html_body = convert_emails_to_links(html_body)
        html_body = process_spacing_tags(html_body)
        
        # Wrap in proper HTML structure
        full_html = f"""
<!DOCTYPE html>
<html>
<head>
    <meta charset="UTF-8">
    <style>
        @page {{
            size: A4;
            margin: 2.5cm;
        }}
        body {{
            font-family: {DEFAULT_FONT_FAMILY};
            font-size: {DEFAULT_FONT_SIZE};
            line-height: {DEFAULT_LINE_HEIGHT};
            color: #000;
        }}
        h1, h2, h3 {{
            color: #000;
            margin-top: 1.5em;
            margin-bottom: 0.6em;
        }}
        h1 {{ font-size: {HEADING_1_SIZE}; font-weight: bold; }}
        h2 {{ font-size: {HEADING_2_SIZE}; font-weight: bold; }}
        p {{
            margin: 0.5em 0;
        }}
        strong {{
            font-weight: bold;
        }}
        ul, ol {{
            margin: 0.5em 0;
            padding-left: 2em;
            line-height: {LIST_LINE_HEIGHT};
        }}
        li {{
            margin: 0.3em 0;
            line-height: {LIST_LINE_HEIGHT};
        }}
        a.survey-link {{
            color: #0066cc;
            text-decoration: underline;
        }}
        a.email-link {{
            color: #0066cc;
            text-decoration: underline;
        }}
    </style>
</head>
<body>
{html_body}
</body>
</html>
"""
        return full_html
        
    except Exception as e:
        raise Exception(f"Failed to load template in worker: {e}")


def convert_emails_to_links(html_content):
    """Convert email addresses in HTML to clickable mailto: links"""
    email_pattern = r'\b([a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,})\b'
    
    def replace_email(match):
        email = match.group(1)
        return f'<a href="mailto:{email}" class="email-link">{email}</a>'
    
    return re.sub(email_pattern, replace_email, html_content)


def process_spacing_tags(html_content):
    """Process spacing tags from Word document"""
    tag_count = html_content.count('&lt;linjeskift&gt;')
    if tag_count > 0:
        html_content = html_content.replace('&lt;linjeskift&gt;', '<p>&nbsp;</p>')
    return html_content


def process_single_file_worker(row_dict, output_folder, required_placeholders):
    """
    Worker function for parallel processing.
    Bruker cached template fra worker process.
    """
    global _worker_template_cache
    
    try:
        # Use cached template
        content = _worker_template_cache
        
        # Replace placeholders
        for placeholder in required_placeholders:
            value = str(row_dict.get(placeholder, ''))
            
            # Special handling for survey link
            if placeholder.lower() == 'uniksurveylink':
                value = f'<a href="{value}" class="survey-link">{value}</a>'
            
            content = content.replace(f'[{placeholder}]', value)
        
        # Generate filename
        filename_tag = str(row_dict['filename_tag'])
        filename_tag_clean = filename_tag.replace('/', '_').replace('\\', '_').replace(':', '_')
        filename = f"brakar_opinion_{filename_tag_clean}.pdf"
        pdf_path = os.path.join(output_folder, filename)
        
        # Create PDF using xhtml2pdf
        with open(pdf_path, 'wb') as pdf_file:
            pisa_status = pisa.CreatePDF(
                content.encode('utf-8'),
                dest=pdf_file,
                encoding='utf-8'
            )
        
        if pisa_status.err:
            return {'success': False, 'error': 'xhtml2pdf conversion failed'}
        
        return {'success': True, 'filename': filename}
    
    except Exception as e:
        filename_tag = row_dict.get('filename_tag', 'unknown')
        return {'success': False, 'error': str(e)}


# =============================================================================
# HOVEDKLASSE
# =============================================================================

class HTMLPDFGenerator:
    """
    Fast HTML-based PDF generator with parallel processing using xhtml2pdf
    
    OPTIMALISERT FOR 60 000 FILER:
    - Worker template caching (reduserer pickle overhead)
    - Kun nødvendige kolonner leses fra Excel
    - Optimalisert multiprocessing
    - FIKSET: Data caching problem - reload ved forskjellig limit_rows
    """
    
    def __init__(self, template_path, lookup_file, output_folder, 
                 required_placeholders, batch_size=2000, num_workers=None):
        """
        Initialize generator
        
        Parameters:
        -----------
        template_path : str or Path
            Path to Word template
        lookup_file : str or Path
            Path to Excel file
        output_folder : str or Path
            Output folder for PDFs
        required_placeholders : list
            List of required placeholder names (without [])
        batch_size : int
            Number of files to process per batch
        num_workers : int
            Number of parallel workers (default: CPU count - 1)
        """
        self.template_path = str(Path(template_path).resolve())
        self.lookup_file = str(Path(lookup_file).resolve())
        self.output_folder = str(Path(output_folder).resolve())
        self.required_placeholders = required_placeholders
        self.batch_size = batch_size
        
        # OPTIMALISERING: Bruk flere workers for bedre ytelse
        self.num_workers = num_workers or max(1, multiprocessing.cpu_count() - 1)
        
        # Create output folder
        Path(self.output_folder).mkdir(parents=True, exist_ok=True)
        
        # Cleanup old temp folders
        self.cleanup_temp_folders()
        
        # Setup logging
        self.setup_logging()
        
        # Data cache - FIKSET: Hold også styr på hvor mange rader som er lastet
        self.df = None
        self._cached_limit_rows = None
        
        # Stats
        self.stats = {
            'total': 0,
            'success': 0,
            'failed': 0,
            'start_time': None,
            'end_time': None
        }
    
    def cleanup_temp_folders(self):
        """Remove old temp folders from previous runs"""
        for folder in Path(self.output_folder).glob("temp_*"):
            if folder.is_dir():
                try:
                    shutil.rmtree(folder)
                except:
                    pass
    
    def setup_logging(self):
        """Setup logging to file and console"""
        log_filename = os.path.join(
            self.output_folder, 
            f'generation_log_{datetime.now().strftime("%Y%m%d_%H%M%S")}.log'
        )
        
        # Clear any existing handlers
        for handler in logging.getLogger().handlers[:]:
            logging.getLogger().removeHandler(handler)
        
        logging.basicConfig(
            level=logging.INFO,
            format='%(asctime)s - %(levelname)s - %(message)s',
            handlers=[
                logging.FileHandler(log_filename, encoding='utf-8'),
                logging.StreamHandler()
            ]
        )
        
        self.logger = logging.getLogger(__name__)
        self.logger.info(f"Log file created: {log_filename}")
    
    def load_data(self, limit_rows=None):
        """
        Load Excel data with optimized column selection.
        
        FIKSET: Reloader data hvis limit_rows er forskjellig fra cached versjon
        
        Parameters:
        -----------
        limit_rows : int, optional
            Limit number of rows to read (for testing)
        """
        # FIKSET: Sjekk om vi må reloade data
        if self.df is not None and self._cached_limit_rows == limit_rows:
            self.logger.info(f"Using cached data ({len(self.df)} rows)")
            return self.df
        
        # Reload hvis limit_rows er forskjellig
        if self.df is not None and self._cached_limit_rows != limit_rows:
            self.logger.info(f"Reloading data: cached={self._cached_limit_rows}, requested={limit_rows}")
        
        self.logger.info("\nReading Excel file...")
        
        # Determine file type (Excel or CSV)
        file_extension = Path(self.lookup_file).suffix.lower()

        # OPTIMALISERING: Les kun nødvendige kolonner for å spare minne
        required_cols = self.required_placeholders + ['filename_tag']
        
        try:
            if file_extension == '.xlsx' or file_extension == '.xls':
                # Read Excel file
                df = pd.read_excel(
                    self.lookup_file,
                    usecols=required_cols,  # Only required columns
                    nrows=limit_rows,       # Limit rows for testing
                    engine='openpyxl',      # Use openpyxl engine
                    dtype={'filename_tag': str}  # Preserve leading zeros
                )
            elif file_extension == '.csv':
                # Read CSV file
                df = pd.read_csv(
                    self.lookup_file,
                    usecols=required_cols,  # Only required columns
                    nrows=limit_rows,       # Limit rows for testing
                    dtype={'filename_tag': str}  # Preserve leading zeros
                )
            else:
                raise ValueError(f"Unsupported file type: {file_extension}")
        except ValueError as e:
            # Hvis usecols feiler, prøv uten
            self.logger.warning(f"Could not limit columns: {e}")
            df = pd.read_excel(
                self.lookup_file, 
                nrows=limit_rows,
                engine='openpyxl',
                dtype={'filename_tag': str}
            )
            # Filtrer til kun nødvendige kolonner
            available_cols = [col for col in required_cols if col in df.columns]
            df = df[available_cols]
        
        self.df = df
        self._cached_limit_rows = limit_rows  # FIKSET: Lagre limit_rows
        self.stats['total'] = len(df)
        self.logger.info(f"  Loaded {self.stats['total']} rows")
        self.logger.info(f"  Columns: {list(df.columns)}")
        
        return df
    
    def validate_excel_columns(self, df):
        """Validate Excel has required columns"""
        missing = []
        required_all = self.required_placeholders + ['filename_tag']
        
        for col in required_all:
            if col not in df.columns:
                missing.append(col)
        
        if missing:
            self.logger.error(f"Missing columns in Excel: {missing}")
            self.logger.info(f"Available columns: {list(df.columns)}")
            return False
        
        return True
    
    def process_batch(self, batch_df, batch_num, output_folder):
        """Process a batch of files in parallel with worker template caching"""
        self.logger.info(f"\nProcessing batch {batch_num} ({len(batch_df)} files)...")
        
        batch_success = 0
        batch_failed = 0
        
        # OPTIMALISERING: Workers initialiseres med template via initializer
        with ProcessPoolExecutor(
            max_workers=self.num_workers,
            initializer=init_worker,
            initargs=(self.template_path,)
        ) as executor:
            # Submit all tasks
            futures = {}
            for idx, row in batch_df.iterrows():
                # OPTIMALISERING: Send kun nødvendig data (ikke hele row)
                row_dict = {
                    col: row[col] 
                    for col in self.required_placeholders + ['filename_tag']
                }
                
                future = executor.submit(
                    process_single_file_worker,
                    row_dict,
                    output_folder,
                    self.required_placeholders
                )
                futures[future] = row.get('filename_tag', idx)
            
            # Collect results
            for future in as_completed(futures):
                filename_tag = futures[future]
                try:
                    result = future.result()
                    if result['success']:
                        batch_success += 1
                        self.stats['success'] += 1
                    else:
                        batch_failed += 1
                        self.stats['failed'] += 1
                        self.logger.error(f"Failed {filename_tag}: {result.get('error', 'Unknown')}")
                except Exception as e:
                    batch_failed += 1
                    self.stats['failed'] += 1
                    self.logger.error(f"Failed {filename_tag}: {str(e)}")
            
            # Progress update
            elapsed = (datetime.now() - self.stats['processing_start_time']).total_seconds()
            rate = self.stats['success'] / elapsed if elapsed > 0 else 0
            remaining = (self.stats['total'] - self.stats['success']) / rate if rate > 0 else 0
            
            self.logger.info(
                f"Progress: {self.stats['success']}/{self.stats['total']} "
                f"({self.stats['success']/self.stats['total']*100:.1f}%) - "
                f"Rate: {rate:.1f} files/sec - ETA: {remaining/60:.1f} min"
            )
        
        self.logger.info(f"Batch {batch_num} complete: {batch_success} success, {batch_failed} failed")
        gc.collect()
        
        return batch_success, batch_failed
    

    def process_batch(self, batch_df, batch_num, output_folder):
        """
        Process a batch of files in parallel with worker template caching
        """
        self.logger.info(f"\nProcessing batch {batch_num} ({len(batch_df)} files)...")
        
        batch_success = 0
        batch_failed = 0
        
        # OPTIMALISERING: Workers initialiseres med template via initializer
        with ProcessPoolExecutor(
            max_workers=self.num_workers,
            initializer=init_worker,
            initargs=(self.template_path,)
        ) as executor:
            # Submit all tasks
            futures = {}
            for idx, row in batch_df.iterrows():
                # OPTIMALISERING: Send kun nødvendig data (ikke hele row)
                row_dict = {
                    col: row[col] 
                    for col in self.required_placeholders + ['filename_tag']
                }
                
                future = executor.submit(
                    process_single_file_worker,
                    row_dict,
                    output_folder,
                    self.required_placeholders
                )
                futures[future] = row.get('filename_tag', idx)
            
            # Collect results with progress bar
            with tqdm(total=len(futures), desc=f"Batch {batch_num}") as pbar:
                for future in as_completed(futures):
                    filename_tag = futures[future]
                    try:
                        result = future.result()
                        if result['success']:
                            batch_success += 1
                            self.stats['success'] += 1
                        else:
                            batch_failed += 1
                            self.stats['failed'] += 1
                            self.logger.error(f"Failed {filename_tag}: {result.get('error', 'Unknown')}")
                    except Exception as e:
                        batch_failed += 1
                        self.stats['failed'] += 1
                        self.logger.error(f"Failed {filename_tag}: {str(e)}")
                    finally:
                        pbar.update(1)  # Oppdater progress bar
            
            self.logger.info(f"Batch {batch_num} complete: {batch_success} success, {batch_failed} failed")
            gc.collect()
            
            return batch_success, batch_failed


    def generate(self, limit_rows=None):
        """
        Main generation process
        
        Parameters:
        -----------
        limit_rows : int, optional
            Limit number of rows to process (for testing).
            Bruk limit_rows=1 for å teste én fil
            Bruk limit_rows=100, 1000, 10000 for større tester
            Bruk limit_rows=None for full produksjon
        """
        # Reset stats for hver generate() kjøring
        self.stats = {
            'total': 0,
            'success': 0,
            'failed': 0,
            'start_time': datetime.now(),
            'processing_start_time': None,
            'end_time': None
        }
        
        self.logger.info(f"\n{'='*60}")
        self.logger.info(f"Starting PDF Generation (xhtml2pdf + worker caching)")
        self.logger.info(f"{'='*60}")
        self.logger.info(f"Template: {self.template_path}")
        self.logger.info(f"Excel: {self.lookup_file}")
        self.logger.info(f"Output: {self.output_folder}")
        self.logger.info(f"Batch size: {self.batch_size}")
        self.logger.info(f"Parallel workers: {self.num_workers}")
        
        if limit_rows:
            self.logger.info(f"TESTING MODE: Limited to {limit_rows} rows")
        
        # Load data
        try:
            df = self.load_data(limit_rows=limit_rows)
        except Exception as e:
            self.logger.error(f"Failed to read Excel: {e}")
            return False
        
        # Validate
        if not self.validate_excel_columns(df):
            return False
        
        # Start processing timer AFTER data is loaded
        self.stats['processing_start_time'] = datetime.now()
        
        # Process in batches
        num_batches = (len(df) + self.batch_size - 1) // self.batch_size
        
        for batch_num in range(num_batches):
            start_idx = batch_num * self.batch_size
            end_idx = min((batch_num + 1) * self.batch_size, len(df))
            batch_df = df.iloc[start_idx:end_idx]
            
            self.process_batch(batch_df, batch_num + 1, self.output_folder)
        
        # Final stats
        self.stats['end_time'] = datetime.now()
        duration = (self.stats['end_time'] - self.stats['processing_start_time']).total_seconds()
        total_duration = (self.stats['end_time'] - self.stats['start_time']).total_seconds()
        
        self.logger.info(f"\n{'='*60}")
        self.logger.info(f"GENERATION COMPLETE")
        self.logger.info(f"{'='*60}")
        self.logger.info(f"Total files: {self.stats['total']}")
        self.logger.info(f"Successful: {self.stats['success']}")
        self.logger.info(f"Failed: {self.stats['failed']}")
        self.logger.info(f"Processing time: {duration/60:.1f} minutes")
        self.logger.info(f"Total time (incl. data loading): {total_duration/60:.1f} minutes")
        self.logger.info(f"Rate: {self.stats['success']/duration:.2f} files/second")
        self.logger.info(f"{'='*60}")
        
        return True


if __name__ == "__main__":
    #template_file = TEMPLATE_FILE
    #lookup_file = LOOKUP_FILE
    #output_folder = OUTPUT_FOLDER
    
    # Required placeholders (without brackets)
    required_placeholders = [
        "mic",
        "navn",
        "uniksurveylink"
    ]
    
    # Create generator
    generator = HTMLPDFGenerator(
        template_path=TEMPLATE_FILE,
        lookup_file=LOOKUP_FILE,
        output_folder=OUTPUT_FOLDER,
        required_placeholders=required_placeholders,
        batch_size=2000,
        num_workers=None  # Auto-detect CPU cores
    )
    
    print("\n" + "="*70)
    print("OPTIMALISERT VERSJON - TESTING WORKFLOW")
    print("="*70)
    
    # Lag testmappe for PDF-er
    test_output_folder = Path(OUTPUT_FOLDER).joinpath("TEST")
    test_output_folder.mkdir(parents=True, exist_ok=True)
    
    # STEG 1: Les inn og behandle de første 20 radene
    print("\nTEST: Genererer PDF-er for de første 20 radene...")
    generator.output_folder = str(test_output_folder)
    success = generator.generate(limit_rows=20)
    
    if success:
        print("\n✅ Test med 20 rader fullført.")
        print(f"PDF-er er lagret i testmappen: {test_output_folder}")
        print("\nVennligst sjekk resultat og bekreft om du ønsker å fortsette.")
        
        # Spør brukeren om neste steg
        while True:
            user_input = input("\nFortsett med hele datasettet? ('Y' eller angi antall): ").strip().lower()
            if user_input == "y":
                print("\n✅ Starter full produksjon...")
                generator.output_folder = str(OUTPUT_FOLDER)
                generator.generate()
                break
            elif user_input.isdigit():
                num_rows = int(user_input)
                print(f"\n✅ Starter produksjon med {num_rows} rader...")
                generator.generate(limit_rows=num_rows)
                break
            else:
                print("Ugyldig valg. Skriv 'Y' eller oppgi et antall rader.")
    
    else:
        print("\n❌ Test med 20 rader feilet. Sjekk loggfilen for detaljer.")
